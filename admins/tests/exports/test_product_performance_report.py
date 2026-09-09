import csv
import io
import secrets
from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from admins.models import ProductCostProfile
from admins.services.product_performance_service import (
    ProductPerformanceError,
    product_performance_report,
)
from base.models import Category, Order, OrderItem, OrderRefund, Product, Session
from base.repositories.session import SessionRepository
from stock.models import StockItem, StockLocation, StockTransaction, StockUnit


pytestmark = pytest.mark.django_db


def _at(day, hour=12):
    return timezone.make_aware(
        datetime.combine(day, datetime.min.time()).replace(hour=hour),
        timezone.get_current_timezone(),
    )


def _sale(user, product, day, *, quantity=1, unit_price='100.00'):
    total = Decimal(unit_price) * quantity
    order = Order.objects.create(
        user=user,
        cashier=user,
        status=Order.Status.COMPLETED,
        is_paid=True,
        paid_at=_at(day),
        payment_method=Order.PaymentMethod.CASH,
        subtotal=total,
        total_amount=total,
        branch_id='branch1',
    )
    item = OrderItem.objects.create(
        order=order,
        product=product,
        quantity=quantity,
        original_price=unit_price,
        price=unit_price,
        branch_id='branch1',
    )
    return order, item


def _actual_cost(user, order, item, amount, suffix='A'):
    unit = StockUnit.objects.create(
        name=f'Piece {suffix}', short_name=f'p{suffix}', unit_type='COUNT',
        is_base_unit=True, branch_id='branch1',
    )
    location = StockLocation.objects.create(
        name=f'Kitchen {suffix}', type='KITCHEN', branch_id='branch1',
    )
    stock_item = StockItem.objects.create(
        name=f'Ingredients {suffix}', base_unit=unit, item_type='RAW',
        avg_cost_price=amount, branch_id='branch1',
    )
    return StockTransaction.objects.create(
        transaction_number=f'PERF-COST-{suffix}-{item.id}',
        stock_item=stock_item,
        location=location,
        movement_type=StockTransaction.MovementType.SALE_OUT,
        quantity=item.quantity,
        unit=unit,
        base_quantity=item.quantity,
        quantity_before='10',
        quantity_after=Decimal('10') - item.quantity,
        unit_cost=Decimal(amount) / item.quantity,
        total_cost=amount,
        order=order,
        order_item=item,
        user=user,
        branch_id='branch1',
    )


def _profile(user, product, cost, start, end=None):
    return ProductCostProfile.objects.create(
        branch_id='branch1',
        product=product,
        treatment=ProductCostProfile.Treatment.STANDARD,
        standard_unit_cost=cost,
        effective_from=start,
        effective_to=end,
        verified_by=user,
        verified_at=timezone.now(),
    )


def _report(**kwargs):
    return product_performance_report(
        branch_id='branch1',
        preset='custom',
        date_from='2026-08-01',
        date_to='2026-08-31',
        paginate=False,
        **kwargs,
    )


def _admin_auth(admin_user):
    token = secrets.token_hex(32)
    Session.objects.create(
        user_id=admin_user,
        ip_address='127.0.0.1',
        payload=SessionRepository.hash_token(token),
        expires_at=timezone.now() + timedelta(hours=1),
    )
    return {'HTTP_AUTHORIZATION': f'Bearer {token}'}


def test_actual_order_line_cost_drives_profit_and_not_current_stock_cost(
    admin_user, product,
):
    order, item = _sale(
        admin_user, product, date(2026, 8, 10), quantity=2, unit_price='100',
    )
    transaction = _actual_cost(admin_user, order, item, '60', 'ACTUAL')
    StockItem.objects.filter(pk=transaction.stock_item_id).update(
        avg_cost_price='9999',
    )

    report = _report()

    assert report['summary']['total_units_sold'] == 2
    assert report['summary']['total_revenue'] == '200.00'
    assert report['summary']['total_ingredient_cost'] == '60.00'
    assert report['summary']['total_gross_profit'] == '140.00'
    row = report['products'][0]
    assert row['selling_price_per_unit'] == '100.0000'
    assert row['ingredient_cost_per_unit'] == '30.0000'
    assert row['gross_profit_per_item'] == '70.0000'
    assert row['gross_profit_margin_pct'] == '70.00'
    assert row['cost_source'] == 'ACTUAL_STOCK'
    assert row['cost_complete'] is True


def test_effective_dated_verified_costs_are_used_for_each_sale(admin_user, product):
    _sale(admin_user, product, date(2026, 8, 10), unit_price='100')
    _sale(admin_user, product, date(2026, 8, 20), unit_price='100')
    _profile(
        admin_user, product, '10', date(2026, 8, 1), date(2026, 8, 15),
    )
    _profile(admin_user, product, '30', date(2026, 8, 16))

    row = _report()['products'][0]

    assert row['units_sold'] == 2
    assert row['ingredient_cost_per_unit'] == '20.0000'
    assert row['total_ingredient_cost'] == '40.00'
    assert row['gross_profit'] == '160.00'
    assert row['cost_source'] == 'VERIFIED_STANDARD'


def test_full_cancellation_reverses_units_revenue_and_historical_cost(
    admin_user, product,
):
    order, item = _sale(admin_user, product, date(2026, 8, 10), unit_price='100')
    _actual_cost(admin_user, order, item, '30', 'REFUND')
    OrderRefund.objects.create(
        order=order,
        amount='100',
        cash_amount='100',
        drawer_cash_amount='100',
        refunded_at=_at(date(2026, 8, 11)),
        source=OrderRefund.Source.ORDER_CANCEL,
        source_id='product-performance-refund',
        branch_id='branch1',
    )

    report = _report()
    row = report['products'][0]

    assert row['units_sold'] == 1
    assert row['units_refunded'] == 1
    assert row['net_units'] == 0
    assert row['total_revenue'] == '0.00'
    assert row['total_ingredient_cost'] == '0.00'
    assert row['gross_profit'] == '0.00'
    assert report['summary']['total_gross_profit'] == '0.00'


def test_discount_and_provider_refund_reduce_revenue_without_returning_cost(
    admin_user, product,
):
    order, item = _sale(
        admin_user, product, date(2026, 8, 10), quantity=2, unit_price='100',
    )
    Order.objects.filter(pk=order.pk).update(
        discount_amount='50', total_amount='150',
    )
    _actual_cost(admin_user, order, item, '60', 'DISCOUNT')
    OrderRefund.objects.create(
        order=order,
        amount='30',
        cash_amount='30',
        drawer_cash_amount='0',
        refunded_at=_at(date(2026, 8, 11)),
        source=OrderRefund.Source.COURIER_PAYMENT,
        source_id='product-performance-provider-refund',
        branch_id='branch1',
    )

    row = _report()['products'][0]

    assert row['selling_price_per_unit'] == '75.0000'
    assert row['units_refunded'] == 0
    assert row['total_revenue'] == '120.00'
    assert row['total_ingredient_cost'] == '60.00'
    assert row['gross_profit'] == '60.00'
    assert row['gross_profit_margin_pct'] == '50.00'


def test_missing_historical_cost_never_becomes_fake_zero_profit(
    admin_user, product,
):
    _sale(admin_user, product, date(2026, 8, 10), unit_price='100')

    report = _report()
    row = report['products'][0]

    assert row['total_ingredient_cost'] is None
    assert row['gross_profit'] is None
    assert row['gross_profit_margin_pct'] is None
    assert row['cost_complete'] is False
    assert report['summary']['total_ingredient_cost'] is None
    assert report['summary']['total_gross_profit'] is None
    assert report['summary']['products_missing_cost'] == 1
    assert report['coverage']['missing_cost_products'][0]['product_id'] == product.id


def test_category_filter_and_profit_sort_apply_before_totals(admin_user, product):
    first_category = product.category
    second_category = Category.objects.create(name='Drinks')
    second = Product.objects.create(
        name='Lemonade', category=second_category, price='200',
    )
    _sale(admin_user, product, date(2026, 8, 10), unit_price='100')
    _sale(admin_user, second, date(2026, 8, 10), unit_price='200')
    _profile(admin_user, product, '90', date(2026, 8, 1))
    _profile(admin_user, second, '20', date(2026, 8, 1))

    all_rows = _report(sort='highest_profit')['products']
    filtered = _report(category_id=first_category.id)

    assert [row['product_name'] for row in all_rows] == ['Lemonade', product.name]
    assert filtered['summary']['product_count'] == 1
    assert filtered['summary']['total_revenue'] == '100.00'
    assert filtered['products'][0]['product_id'] == product.id


def test_presets_and_invalid_ranges_use_business_dates(admin_user):
    now = _at(date(2026, 9, 10), 12)
    last_seven = product_performance_report(
        branch_id='branch1', preset='last_7_days', now=now,
    )
    last_month = product_performance_report(
        branch_id='branch1', preset='last_month', now=now,
    )
    assert last_seven['range']['from'] == '2026-09-04'
    assert last_seven['range']['to'] == '2026-09-10'
    assert last_month['range']['from'] == '2026-08-01'
    assert last_month['range']['to'] == '2026-08-31'

    with pytest.raises(ProductPerformanceError) as exc_info:
        product_performance_report(
            branch_id='branch1', preset='custom',
            date_from='2026-08-31', date_to='2026-08-01', now=now,
        )
    assert exc_info.value.code == 'INVALID_REPORT_RANGE'


def test_json_and_all_export_formats_share_the_same_complete_dataset(
    admin_user, product,
):
    product.name = '=SUM(1,1)'
    product.save(update_fields=['name'])
    order, item = _sale(admin_user, product, date(2026, 8, 10), unit_price='100')
    _actual_cost(admin_user, order, item, '25', 'EXPORT')
    auth = _admin_auth(admin_user)
    client = Client()
    query = 'preset=custom&from=2026-08-01&to=2026-08-31&sort=highest_profit'

    json_response = client.get(
        f'/api/admins/reports/product-performance?{query}', **auth,
    )
    assert json_response.status_code == 200, json_response.content
    assert json_response.json()['data']['summary']['total_gross_profit'] == '75.00'

    xlsx = client.get(
        f'/api/admins/reports/product-performance/export?{query}&format=xlsx',
        **auth,
    )
    assert xlsx.status_code == 200, xlsx.content
    assert xlsx['X-Export-Count'] == '1'
    assert xlsx['X-Report-Cost-Complete'] == 'true'
    workbook = load_workbook(io.BytesIO(xlsx.content), data_only=False)
    assert workbook.sheetnames == ['Summary', 'Products', 'Categories', 'Daily', 'Methodology']
    assert workbook['Products']['C5'].value == "'=SUM(1,1)"
    assert workbook['Products']['U5'].value == 75

    csv_response = client.get(
        f'/api/admins/reports/product-performance/export?{query}&format=csv',
        **auth,
    )
    assert csv_response.status_code == 200
    decoded = csv_response.content.decode('utf-8-sig')
    parsed = list(csv.reader(io.StringIO(decoded)))
    assert any("'=SUM(1,1)" in row for row in parsed)
    assert any(row and row[0] == 'TOTAL' and row[15] == '75.00' for row in parsed)

    pdf = client.get(
        f'/api/admins/reports/product-performance/export?{query}&format=pdf',
        **auth,
    )
    assert pdf.status_code == 200, pdf.content[:500]
    assert pdf.content.startswith(b'%PDF-')
    assert len(pdf.content) > 5_000


def test_report_is_manager_admin_only_and_export_format_is_validated(
    client, cashier_user, admin_user,
):
    from base.models import User

    path = '/api/admins/reports/product-performance?preset=today'
    assert client.get(path).status_code == 401

    cashier_auth = _admin_auth(cashier_user)
    assert Client().get(path, **cashier_auth).status_code == 403

    manager = User.objects.create(
        email='report-manager@example.com',
        first_name='Report',
        last_name='Manager',
        password='not-used-by-session-test',
        role=User.RoleChoices.MANAGER,
        status=User.UserStatus.ACTIVE,
    )
    manager_auth = _admin_auth(manager)
    manager_response = Client().get(
        f'{path}&branch_id=other-restaurant', **manager_auth,
    )
    assert manager_response.status_code == 200
    assert manager_response.json()['data']['branch_id'] == 'branch1'

    admin_auth = _admin_auth(admin_user)
    invalid = Client().get(
        '/api/admins/reports/product-performance/export?preset=today&format=docx',
        **admin_auth,
    )
    assert invalid.status_code == 422
    assert invalid.json()['code'] == 'INVALID_EXPORT_FORMAT'
